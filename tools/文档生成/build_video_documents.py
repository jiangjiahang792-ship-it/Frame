"""根据现有操作录屏生成中文 Word 操作流程文档。"""

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


# 项目根目录，用于统一解析输入素材和输出文档。
PROJECT_ROOT = Path(r"D:\MyCode\PublicWook\TDJS-Vision")
# 五秒抽帧目录，保存从现有 MP4 中提取的原始画面。
FRAME_DIRECTORY = PROJECT_ROOT / "artifacts" / "video_docs_work" / "frames_5s"
# 文档专用图片目录，保存裁剪后的关键画面和转换后的图标。
PREPARED_IMAGE_DIRECTORY = PROJECT_ROOT / "artifacts" / "video_docs_work" / "prepared"
# 工作区内的最终文档目录，桌面交付前先在这里完成渲染验收。
OUTPUT_DIRECTORY = PROJECT_ROOT / "artifacts" / "操作流程文档"
# 用户指定的 ICO 图标路径。
ICON_PATH = PROJECT_ROOT / "ig6s1-ksbbd-001.ico"
# 现有视频路径，仅对实际存在的 MP4 生成文档。
VIDEO_PATH = Path(r"C:\Users\34652\Desktop\框架操作流程视频\搭建AI检测端子线芯视频.mp4")

# 版式主色，对应 compact_reference_guide 的标题蓝色。
HEADING_BLUE = "2E74B5"
# 次级标题颜色，对应 compact_reference_guide 的深蓝色。
HEADING_DARK_BLUE = "1F4D78"
# 正文深灰色，保证打印和屏幕阅读清晰。
BODY_COLOR = "222222"
# 辅助信息灰色，用于时间、来源和页脚。
MUTED_COLOR = "666666"
# 提示底色，用于强调容易出错的配置项。
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, latin_font="Calibri", east_asia_font="Microsoft YaHei", size=None,
                 color=BODY_COLOR, bold=None, italic=None):
    """为文字片段显式设置中西文字体、字号、颜色和强调样式。"""

    run.font.name = latin_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document_styles(document):
    """按照紧凑型操作指南预设配置正文、标题、图注和提示样式。"""

    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.color.rgb = RGBColor.from_string(BODY_COLOR)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(HEADING_DARK_BLUE)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED_COLOR)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 18, 10),
        "Heading 2": (13, HEADING_BLUE, 14, 7),
        "Heading 3": (12, HEADING_DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(MUTED_COLOR)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "提示" not in styles:
        callout = styles.add_style("提示", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["提示"]
    callout.base_style = normal
    callout.font.name = "Calibri"
    callout._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    callout._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    callout.font.size = Pt(10.5)
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.line_spacing = 1.2
    style_ppr = callout._element.get_or_add_pPr()
    shading = style_ppr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        style_ppr.append(shading)
    shading.set(qn("w:fill"), CALLOUT_FILL)


def configure_page_geometry(document):
    """显式设置 Letter 纸张、页边距和页眉页脚距离。"""

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True


def add_page_number(paragraph):
    """向页脚段落插入可由 Word 和 LibreOffice更新的 PAGE 字段。"""

    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, size=9, color=MUTED_COLOR)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run = OxmlElement("w:r")
    field_run.append(begin)
    field_run.append(instruction)
    field_run.append(separate)
    field_run.append(result)
    field_run.append(end)
    paragraph._p.append(field_run)
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED_COLOR)


def configure_header_and_footer(document, logo_path):
    """在正文页设置安静的运行页眉和右对齐页码。"""

    section = document.sections[0]
    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    logo_run = header_paragraph.add_run()
    logo_run.add_picture(str(logo_path), width=Inches(0.22))
    label_run = header_paragraph.add_run("  TDJS 视觉操作指南")
    set_run_font(label_run, size=9, color=MUTED_COLOR, bold=True)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.paragraph_format.space_before = Pt(0)
    add_page_number(footer_paragraph)


def create_numbering_definition(document, number_format, level_text, font_name="Calibri"):
    """创建符合紧凑型指南缩进和行距要求的真实 Word 编号定义。"""

    numbering = document.part.numbering_part.element
    abstract_ids = [int(element.get(qn("w:abstractNumId"))) for element in numbering.findall(qn("w:abstractNum"))]
    number_ids = [int(element.get(qn("w:numId"))) for element in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    number_id = max(number_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level_type)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), number_format)
    level.append(num_format)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), level_text)
    level.append(text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)

    paragraph_properties = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    paragraph_properties.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    paragraph_properties.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.append(spacing)
    level.append(paragraph_properties)

    run_properties = OxmlElement("w:rPr")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), font_name)
    run_fonts.set(qn("w:hAnsi"), font_name)
    run_fonts.set(qn("w:eastAsia"), font_name)
    run_properties.append(run_fonts)
    level.append(run_properties)
    abstract.append(level)
    first_number = numbering.find(qn("w:num"))
    if first_number is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_number), abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(number_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return number_id


def add_numbered_item(document, text, number_id):
    """添加使用指定真实编号定义的流程概览项。"""

    paragraph = document.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    number_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    reference = OxmlElement("w:numId")
    reference.set(qn("w:val"), str(number_id))
    number_properties.append(level)
    number_properties.append(reference)
    properties.append(number_properties)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)
    return paragraph


def add_check_item(document, text, check_number_id):
    """添加使用方框符号的真实 Word 检查清单项。"""

    return add_numbered_item(document, text, check_number_id)


def add_labeled_paragraph(document, label, text):
    """添加带加粗中文标签的说明段落。"""

    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=11, color=HEADING_DARK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11)
    return paragraph


def add_callout(document, label, text):
    """添加用于注意事项、故障定位或路径提示的浅色提示段落。"""

    paragraph = document.add_paragraph(style="提示")
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=10.5, color=HEADING_DARK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10.5)
    return paragraph


def prepare_logo():
    """把用户指定的 ICO 图标转换为 Word 可稳定使用的 PNG。"""

    PREPARED_IMAGE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    output_path = PREPARED_IMAGE_DIRECTORY / "tdjs_logo.png"
    with Image.open(ICON_PATH) as image:
        converted = image.convert("RGBA")
        converted.save(output_path)
    return output_path


def frame_path(milliseconds):
    """根据五秒时间点计算抽帧文件路径。"""

    index = milliseconds // 5000 + 1
    return FRAME_DIRECTORY / f"frame-{index:04d}-{milliseconds:06d}ms.jpg"


def prepare_screenshot(milliseconds, output_name):
    """裁掉录屏顶部控制条和底部远程控制浮层，保留主要操作区域。"""

    PREPARED_IMAGE_DIRECTORY.mkdir(parents=True, exist_ok=True)
    source_path = frame_path(milliseconds)
    output_path = PREPARED_IMAGE_DIRECTORY / output_name
    with Image.open(source_path) as image:
        width, height = image.size
        top = min(36, max(0, height - 1))
        bottom = min(835, height)
        if bottom <= top:
            top = 0
            bottom = height
        cropped = image.crop((0, top, width, bottom)).convert("RGB")
        cropped.save(output_path, quality=94)
    return output_path


def add_screenshot(document, image_path, caption_text, alt_text, width_inches=6.5):
    """插入全宽操作截图、无障碍替代文字和居中图注。"""

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    picture = run.add_picture(str(image_path), width=Inches(width_inches))
    picture._inline.docPr.set("descr", alt_text)
    caption = document.add_paragraph(caption_text, style="Caption")
    caption.paragraph_format.keep_together = True
    return caption


def add_cover_page(document, logo_path):
    """创建简洁的操作指南封面和来源说明。"""

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_run = logo_paragraph.add_run()
    logo_run.add_picture(str(logo_path), width=Inches(0.48))
    logo_paragraph.paragraph_format.space_after = Pt(14)

    title = document.add_paragraph("搭建 AI 检测端子线芯操作流程", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("检测 1、检测 2 双流程图文操作指南", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_before = Pt(16)
    source.paragraph_format.space_after = Pt(6)
    source_run = source.add_run("对应视频：搭建AI检测端子线芯视频.mp4")
    set_run_font(source_run, size=10.5, color=MUTED_COLOR)

    duration = document.add_paragraph()
    duration.alignment = WD_ALIGN_PARAGRAPH.CENTER
    duration.paragraph_format.space_after = Pt(6)
    duration_run = duration.add_run("视频时长：08:45  |  软件：机器视觉AI检测系统 V1.0.0.0")
    set_run_font(duration_run, size=10.5, color=MUTED_COLOR)

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(28)
    date_run = date.add_run("整理日期：2026 年 7 月 14 日")
    set_run_font(date_run, size=10, color=MUTED_COLOR)
    document.add_page_break()


def add_overview(document, decimal_number_id):
    """添加目标、流程结构、准备项和操作步骤概览。"""

    document.add_heading("一、文档目标", level=1)
    document.add_paragraph(
        "本指南把录屏中的完整操作整理为可复用步骤，用于从零搭建端子线芯 AI 检测流程，"
        "完成图像采集、AI 推理、条件判断、检测结果绘制、图像显示和相机 IO 输出。"
    )

    document.add_heading("二、流程结构", level=1)
    structure = document.add_paragraph()
    structure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    structure.paragraph_format.space_before = Pt(6)
    structure.paragraph_format.space_after = Pt(10)
    run = structure.add_run("图像源 → AI检测 → 多条件判断 → T/F 相机IO分支\n"
                            "多条件判断 → ROI结果绘制 → 图像显示")
    set_run_font(run, size=12, color=HEADING_DARK_BLUE, bold=True)

    document.add_heading("三、开始前准备", level=1)
    add_labeled_paragraph(document, "相机：", "目标工业相机已联网，品牌、序列号和采集画面可正常识别。")
    add_labeled_paragraph(document, "检测项：", "已经准备检测项配置，或可以从现有项目导入检测项。")
    add_labeled_paragraph(document, "模型：", "已经准备加密 AI 配置文件（.ai）和加密模型文件（.engine.encrypted）。")
    add_labeled_paragraph(document, "硬件：", "CPU 或 GPU 推理环境可用；使用 GPU 时确认显卡驱动与模型匹配。")

    document.add_heading("四、操作步骤概览", level=1)
    overview_items = [
        "添加并验证目标相机，进入流程编辑并配置图像源。",
        "添加 AI 检测、多条件判断和双路相机 IO。",
        "配置 ROI 结果绘制与图像显示，初次运行并定位未配置项。",
        "导入检测项并制作 AI 模型配置。",
        "在 AI 节点绑定配置，通过日志确认模型加载。",
        "设置 IO 线路和线路模式。",
        "新建检测 2 流程，复用结构并完成最终验收。",
    ]
    for item in overview_items:
        add_numbered_item(document, item, decimal_number_id)


def add_step(document, number, title, actions, checks, image_specs, callout=None, page_break=True):
    """添加一个包含操作、检查、截图和注意事项的完整步骤页。"""

    document.add_heading(f"步骤 {number}：{title}", level=1)
    for label, text in actions:
        add_labeled_paragraph(document, label, text)
    for image_spec in image_specs:
        image_path, caption, alt_text = image_spec[:3]
        width_inches = image_spec[3] if len(image_spec) > 3 else 6.5
        add_screenshot(document, image_path, caption, alt_text, width_inches)
    if callout is not None:
        add_callout(document, callout[0], callout[1])
    if checks:
        check = document.add_paragraph()
        check_run = check.add_run("本步完成标志：")
        set_run_font(check_run, size=10.5, color=HEADING_DARK_BLUE, bold=True)
        check_text = check.add_run(checks)
        set_run_font(check_text, size=10.5)
    if page_break:
        document.add_page_break()


def build_document():
    """准备关键截图，生成并保存现有视频对应的 Word 操作指南。"""

    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    logo_path = prepare_logo()
    screenshots = {
        "camera": prepare_screenshot(70000, "01_添加目标相机.jpg"),
        "image_source": prepare_screenshot(95000, "02_配置图像源.jpg"),
        "main_flow": prepare_screenshot(140000, "03_主流程与IO分支.jpg"),
        "overlay": prepare_screenshot(170000, "04_ROI绘制与图像显示.jpg"),
        "error": prepare_screenshot(190000, "05_红色节点报错.jpg"),
        "detection_item": prepare_screenshot(295000, "06_固定检测项.jpg"),
        "ai_file": prepare_screenshot(330000, "07_选择AI配置文件.jpg"),
        "ai_editor": prepare_screenshot(340000, "08_AI配置关键字段.jpg"),
        "ai_binding": prepare_screenshot(390000, "09_AI节点绑定配置.jpg"),
        "log": prepare_screenshot(410000, "10_日志确认模型加载.jpg"),
        "camera_io": prepare_screenshot(465000, "11_相机IO配置.jpg"),
        "flow_two": prepare_screenshot(500000, "12_检测2流程.jpg"),
        "final": prepare_screenshot(520000, "13_最终检测结果.jpg"),
    }

    document = Document()
    configure_page_geometry(document)
    configure_document_styles(document)
    configure_header_and_footer(document, logo_path)
    document.core_properties.title = "搭建 AI 检测端子线芯操作流程"
    document.core_properties.subject = "检测 1、检测 2 双流程图文操作指南"
    document.core_properties.author = "TDJS-Vision"
    document.core_properties.keywords = "机器视觉,AI检测,端子线芯,操作流程"
    document.core_properties.comments = "根据搭建AI检测端子线芯视频整理"

    decimal_number_id = create_numbering_definition(document, "decimal", "%1.")
    check_number_id = create_numbering_definition(document, "bullet", "☐", "Segoe UI Symbol")

    add_cover_page(document, logo_path)
    add_overview(document, decimal_number_id)

    add_step(
        document,
        1,
        "添加并验证目标相机",
        [
            ("打开：", "在主界面进入相机管理，点击相机列表旁的“+”新增设备。"),
            ("选择：", "按实际硬件选择相机品牌，搜索设备并确认序列号；输入便于识别的中文自定义名称后添加。"),
            ("验证：", "在添加前先用相机厂商工具确认设备在线、能连续采集，并检查网络与曝光状态。"),
        ],
        "相机出现在相机列表中，实时预览能够稳定显示端子线芯画面。",
        [(screenshots["camera"], "图 1：在相机管理中搜索并添加目标相机（视频 01:10）", "相机管理窗口中选择海康相机并搜索设备")],
        ("注意：", "自定义名称会在图像源和相机 IO 节点中使用，建议按工位或端子位置命名，避免后续选错设备。"),
    )

    add_step(
        document,
        2,
        "进入流程编辑并配置图像源",
        [
            ("进入：", "点击主界面工具栏的“流程管理”，新建或打开“流程1”。"),
            ("添加：", "从左侧“图像采集”分类拖入“图像源”节点，并把它设为流程起点。"),
            ("配置：", "图像源选择“相机”，绑定步骤 1 新增的相机；按现场触发方式设置软触发或硬触发，并确认曝光、增益、触发延迟和超时时间。"),
        ],
        "图像源节点保存后不再显示“未设置”，单步运行可以获得输出图像。",
        [(screenshots["image_source"], "图 2：为图像源选择相机并确认触发、曝光和增益（视频 01:35）", "图像源参数窗口中选择前端端子相机")],
        ("建议：", "调试阶段优先使用软触发；切换到产线硬触发前，再核对触发沿、延迟和相机接线。"),
    )

    add_step(
        document,
        3,
        "搭建 AI 检测和双路 IO 主链",
        [
            ("添加节点：", "依次添加“AI检测”“多条件判断”和两个“相机IO”节点。"),
            ("连接主链：", "按“图像源 → AI检测 → 多条件判断”连接。"),
            ("连接分支：", "把多条件判断的 T、F 两个出口分别连接到两个相机 IO 节点，用于输出不同的 OK/NG 或放行/剔除信号。"),
        ],
        "流程中形成一条 AI 主链和两条条件分支，节点连线方向清晰且无断线。",
        [(screenshots["main_flow"], "图 3：AI 检测、多条件判断和 T/F 相机 IO 分支（视频 02:20）", "流程编辑器中的图像源、AI检测、多条件判断和双路相机IO")],
        ("注意：", "T、F 分支的业务含义必须与现场约定一致。配置 IO 前先明确哪一支代表 OK，哪一支代表 NG。"),
    )

    add_step(
        document,
        4,
        "配置 ROI 结果绘制和图像显示",
        [
            ("结果绘制：", "在“结果处理”分类中添加“ROI结果绘制”，选择图像源输出作为输入图像。"),
            ("订阅结果：", "按需要添加文本、线、矩形或区域；订阅 AI 检测输出，并设置 OK/NG 颜色、字号、线宽、位置和坐标系。"),
            ("图像显示：", "添加“图像显示”节点，订阅 ROI 结果绘制的输出图像，并选择需要展示的图像窗口。"),
        ],
        "运行后，检测文字和框线叠加在采集画面上，并显示到指定图像窗口。",
        [(screenshots["overlay"], "图 4：ROI 结果绘制订阅 AI 结果并输出到图像显示（视频 02:50）", "ROI结果绘制窗口配置文本和矩形叠加")],
        ("显示规则：", "视频中使用绿色显示正常结果、红色显示异常结果。可按项目标准调整，但同一项目应保持一致。"),
    )

    add_step(
        document,
        5,
        "初次运行并定位红色报错节点",
        [
            ("保存：", "保存当前流程后执行一次。"),
            ("观察：", "绿色节点表示执行成功；红色节点表示出现致命错误，需要先处理。"),
            ("定位：", "打开程序运行日志，优先查看“致命”或“异常”分类。此阶段 AI 节点通常因为模型配置或检测项配置未设置而失败。"),
        ],
        "能够明确报错节点和缺失配置，不再盲目修改后续显示或 IO 节点。",
        [(screenshots["error"], "图 5：红色 AI 节点表示致命错误，应先查看日志（视频 03:10）", "流程运行后AI检测和多条件判断节点显示红色")],
        ("常见错误：", "若日志提示“相机对象为空”，返回图像源重新选择相机；若 AI 节点未保存参数，继续完成步骤 6 至步骤 8。"),
    )

    add_step(
        document,
        6,
        "导入并选择检测项配置",
        [
            ("打开：", "回到主界面，依次进入“设置 → 检测项配置”。"),
            ("导入：", "选择此前已经制作好的检测项配置文件并添加。"),
            ("确认：", "返回 AI 检测节点，在“检测项配置”中选择“固定检测项”，确认下拉列表出现刚导入的检测项。"),
        ],
        "AI 节点的固定检测项列表中出现目标工位检测项。",
        [(screenshots["detection_item"], "图 6：固定检测项列表出现新导入的工位检测项（视频 04:55）", "AI检测节点中选择固定检测项和工位检测项")],
        ("注意：", "检测项名称应与当前相机和工位对应。检测 1、检测 2 使用不同配置时，不要复用错误的检测项。"),
    )

    add_step(
        document,
        7,
        "制作 AI 模型配置文件",
        [
            ("打开工具：", "进入“设置 → AI配置工具”，打开一个加密 AI 配置文件（.ai）。"),
            ("选择文件：", "视频示例选择“1-1_encrypted.ai”；实际项目按检测类别选择对应文件。"),
            ("核对字段：", "重点确认 ModelPath、DeviceType 和 ModelType。ModelPath 指向真实的 .engine.encrypted 模型；DeviceType 选择 GPU 或 CPU；ModelType 按模型选择 OBB、DET、SEG 等。"),
            ("保存：", "保存配置文件，并保持加密模型文件位置稳定。"),
        ],
        "",
        [
            (screenshots["ai_file"], "图 7：选择与检测类别对应的加密 AI 配置文件（视频 05:30）", "文件选择窗口中选择加密AI配置文件", 4.8),
            (screenshots["ai_editor"], "图 8：AI 配置中需要重点核对的模型字段（视频 05:40）", "AI配置编辑工具中显示ModelPath、DeviceType和ModelType", 4.8),
        ],
        ("不要照抄：", "视频中的 D 盘路径和显卡型号只是录制环境。ModelPath 必须改成当前电脑的真实绝对路径，否则日志会提示模型加载失败。完成标志：AI 配置文件已保存，模型路径、运行设备和模型类型与当前环境一致。"),
    )

    add_step(
        document,
        8,
        "在 AI 节点绑定模型和检测项",
        [
            ("输入图像：", "AI 检测节点的输入图像选择步骤 2 的图像源输出。"),
            ("模型配置：", "点击“选择”，绑定步骤 7 保存的 AI 配置文件路径。"),
            ("检测项配置：", "选择固定检测项，并选择步骤 6 导入的工位检测项；最后点击“保存”。"),
        ],
        "AI 节点不再显示模型配置为空，固定检测项和模型名称均可见。",
        [(screenshots["ai_binding"], "图 9：在 AI 检测节点中选择 AI 配置文件和固定检测项（视频 06:30）", "AI检测节点参数窗口中绑定模型配置和检测项")],
        ("核对：", "检测 1 和检测 2 的 AI 节点需要分别绑定对应配置；不要只修改一个流程后直接假定另一个流程已同步。"),
    )

    add_step(
        document,
        9,
        "通过日志确认模型加载并验证检测",
        [
            ("查看日志：", "点击右侧“程序运行日志”，搜索“开始加载模型”和“加载完成”。"),
            ("确认环境：", "使用 GPU 时，日志应显示 GPU 运行环境初始化和实际显卡信息。"),
            ("再次运行：", "执行流程，确认图像源、AI 检测、多条件判断、ROI 结果绘制和图像显示依次成功。"),
            ("核对画面：", "图像窗口应显示检测框、测量值和最终 OK/NG 结果。"),
        ],
        "日志明确显示加密模型加载完成，AI 节点为绿色，检测结果正确绘制。",
        [(screenshots["log"], "图 10：底部日志显示加密模型加载完成（视频 06:50）", "程序运行日志中显示engine加密模型加载完成")],
        ("故障定位：", "模型加载失败时依次检查文件是否存在、ModelPath 是否正确、DeviceType 是否匹配硬件、ModelType 是否与模型一致。"),
    )

    add_step(
        document,
        10,
        "配置 T/F 双分支相机 IO",
        [
            ("打开节点：", "分别打开多条件判断 T、F 分支后的两个“相机IO”节点。"),
            ("选择相机：", "选择与当前工位对应的相机。"),
            ("选择模式：", "设置输入/输出模式、实际线路和线路模式；按现场要求设置输出保持时间，并决定是否异步执行。"),
            ("保存测试：", "保存后分别制造 OK、NG 条件，确认两条分支只触发各自对应的 IO。"),
        ],
        "T/F 分支能够输出到正确线路，保持时间和异步设置符合现场节拍。",
        [(screenshots["camera_io"], "图 11：选择相机 IO 的线路、线路模式和保持时间（视频 07:45）", "相机IO设置窗口中选择相机、输入输出、线路和保持时间")],
        ("安全：", "首次联机测试前先断开执行机构或使用安全模式验证信号，确认分支含义和接线无误后再接入产线。"),
    )

    add_step(
        document,
        11,
        "新建检测 2 并复用流程结构",
        [
            ("新建：", "点击流程标签旁的“+”新增“流程2”。"),
            ("复用：", "按检测 1 的结构重新搭建或复制节点：图像源、AI 检测、多条件判断、双路相机 IO、ROI 结果绘制和图像显示。"),
            ("替换配置：", "为检测 2 选择对应相机、AI 配置文件、固定检测项、IO 线路和图像窗口。"),
            ("检查编号：", "复制后节点编号会变化，但连接关系和业务含义应保持一致。"),
        ],
        "流程 1 和流程 2 均包含完整七节点结构，且各自绑定正确配置。",
        [(screenshots["flow_two"], "图 12：检测 2 复用完整节点结构并选择图像窗口（视频 08:20）", "流程2中复用图像源、AI检测、多条件判断、ROI绘制和图像显示")],
        ("注意：", "复制流程后最容易遗漏的是图像窗口和 IO 线路。务必逐个打开节点核对，不要只看连线外观。"),
    )

    add_step(
        document,
        12,
        "运行并完成最终验收",
        [
            ("运行：", "保存方案，分别启动检测 1 和检测 2。"),
            ("检查节点：", "流程状态灯应为绿色，关键节点耗时合理，不应持续出现红色致命错误。"),
            ("检查结果：", "图像窗口显示端子线芯画面、检测框、测量值和 OK/NG 结论。"),
            ("整理窗口：", "把检测 1、检测 2 及辅助窗口调整到便于观察的布局。"),
        ],
        "两条流程均可稳定运行，结果显示正确，日志无未处理致命错误，IO 输出与结果一致。",
        [(screenshots["final"], "图 13：检测 1、检测 2 的结果窗口完成布局（视频 08:40）", "最终界面中多个图像窗口显示端子线芯检测结果")],
        ("性能：", "上线前记录连续运行时的节点耗时和总节拍；若波动明显，优先检查相机超时、模型设备类型和图像显示数量。"),
        page_break=True,
    )

    document.add_heading("最终检查清单", level=1)
    checklist = [
        "目标相机在线，预览和触发稳定。",
        "流程 1、流程 2 的图像源均选择正确相机。",
        "AI 配置中的 ModelPath、DeviceType、ModelType 正确。",
        "固定检测项与当前工位和检测类别一致。",
        "日志显示加密模型加载完成，且无未处理致命错误。",
        "ROI 文本、框线、颜色和图像窗口显示正确。",
        "T/F 分支含义、相机 IO 线路和线路模式已实际验证。",
        "方案已保存，并完成连续运行节拍检查。",
    ]
    for item in checklist:
        add_check_item(document, item, check_number_id)
    add_callout(document, "交付建议：", "完成验收后，把 AI 配置文件、检测项配置、加密模型和方案文件按项目版本一起归档，避免后续只复制方案而遗漏模型。")

    output_path = OUTPUT_DIRECTORY / "搭建AI检测端子线芯操作流程.docx"
    document.save(output_path)
    return output_path


def main():
    """执行文档生成并输出成品路径。"""

    if not VIDEO_PATH.exists():
        raise FileNotFoundError(f"视频文件不存在：{VIDEO_PATH}")
    output_path = build_document()
    print(output_path)


if __name__ == "__main__":
    main()
