"""为新增操作流程视频批量生成中文 Word 操作指南。"""

from dataclasses import dataclass, field
from importlib import util
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


# 项目根目录，用于解析脚本、图标、截图和文档输出路径。
PROJECT_ROOT = Path(r"D:\MyCode\PublicWook\TDJS-Vision")
# 本次新增视频抽帧工作目录，每个视频拥有独立的 frames_5s 与 prepared 子目录。
BATCH_ROOT = PROJECT_ROOT / "artifacts" / "video_docs_work" / "batch_20260714"
# 最终 Word 文档先输出到仓库工作区，完成验收后再复制到用户桌面目录。
OUTPUT_DIRECTORY = PROJECT_ROOT / "artifacts" / "操作流程文档"
# 上一份文档生成脚本中已经包含统一版式、页眉页脚和列表辅助函数，这里只复用公共能力。
BASE_SCRIPT_PATH = PROJECT_ROOT / "tools" / "文档生成" / "build_video_documents.py"
# 用户指定的视频目录，用于在封面和元数据中写入真实来源。
VIDEO_DIRECTORY = Path(r"C:\Users\34652\Desktop\框架操作流程视频")
# 文档整理日期，保持与本批任务日期一致。
DOCUMENT_DATE_TEXT = "2026 年 7 月 14 日"


@dataclass(frozen=True)
class ScreenshotSpec:
    """描述一张要插入文档的关键视频截图。"""

    # 视频时间点，单位毫秒，必须与 5 秒抽帧文件名对应。
    milliseconds: int
    # 输出到 prepared 目录的图片文件名。
    filename: str
    # 图片下方显示的中文图注。
    caption: str
    # Word 图片无障碍替代文字。
    alt_text: str
    # 截图顶部裁剪像素，用于移除录屏控制条或窗口边缘。
    crop_top: int = 0
    # 截图底部裁剪像素，None 表示保留到底部。
    crop_bottom: int | None = None
    # 插入 Word 时使用的图片宽度，单位英寸。
    width_inches: float = 6.35


@dataclass(frozen=True)
class GuideStep:
    """描述文档中的一个操作步骤页面。"""

    # 步骤标题，脚本会自动加上“步骤 N”。
    title: str
    # 操作说明，格式为中文标签和正文内容。
    actions: list[tuple[str, str]]
    # 本步骤完成后应看到的状态。
    completion: str
    # 本步骤对应的一张或多张视频截图。
    screenshots: list[ScreenshotSpec]
    # 重要注意事项，格式为中文标签和正文内容。
    callout: tuple[str, str] | None = None


@dataclass(frozen=True)
class GuideSpec:
    """描述一份完整 Word 操作指南的内容和素材。"""

    # 视频基础名称，同时也是 batch 子目录名。
    video_stem: str
    # 源视频文件名。
    video_file_name: str
    # 文档主标题。
    title: str
    # 文档副标题。
    subtitle: str
    # 输出 Word 文件名。
    output_file_name: str
    # 视频时长说明。
    duration_text: str
    # 文档目标说明。
    goal: str
    # 流程结构或功能结构，用箭头或换行组织。
    structure: str
    # 开始前准备项。
    preparations: list[tuple[str, str]]
    # 操作步骤列表。
    steps: list[GuideStep]
    # 最终检查清单。
    checklist: list[str]
    # 元数据关键字。
    keywords: str = "机器视觉,操作流程,TDJS"


def load_base_helpers():
    """动态加载上一份文档脚本中的版式和组件辅助函数。"""

    module_spec = util.spec_from_file_location("video_document_base", BASE_SCRIPT_PATH)
    if module_spec is None or module_spec.loader is None:
        raise RuntimeError(f"无法加载基础文档脚本：{BASE_SCRIPT_PATH}")
    module = util.module_from_spec(module_spec)
    module_spec.loader.exec_module(module)
    return module


def frame_path(spec: GuideSpec, milliseconds: int) -> Path:
    """根据视频名称和毫秒时间点定位已经抽取出的 JPG 原始帧。"""

    index = milliseconds // 5000 + 1
    return BATCH_ROOT / spec.video_stem / "frames_5s" / f"frame-{index:04d}-{milliseconds:06d}ms.jpg"


def prepare_screenshot(spec: GuideSpec, screenshot: ScreenshotSpec) -> Path:
    """把原始视频帧裁剪为适合 Word 的操作截图。"""

    source_path = frame_path(spec, screenshot.milliseconds)
    if not source_path.exists():
        raise FileNotFoundError(f"缺少视频抽帧文件：{source_path}")

    prepared_directory = BATCH_ROOT / spec.video_stem / "prepared"
    prepared_directory.mkdir(parents=True, exist_ok=True)
    output_path = prepared_directory / screenshot.filename

    with Image.open(source_path) as image:
        width, height = image.size
        top = max(0, min(screenshot.crop_top, height - 1))
        bottom = screenshot.crop_bottom if screenshot.crop_bottom is not None else height
        bottom = max(top + 1, min(bottom, height))
        cropped = image.crop((0, top, width, bottom)).convert("RGB")
        cropped.save(output_path, quality=94)
    return output_path


def add_cover_page(base, document: Document, guide: GuideSpec, logo_path: Path) -> None:
    """创建与既有文档风格一致的封面页。"""

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(52)

    logo_paragraph = document.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_run = logo_paragraph.add_run()
    logo_run.add_picture(str(logo_path), width=Inches(0.48))
    logo_paragraph.paragraph_format.space_after = Pt(14)

    title = document.add_paragraph(guide.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(guide.subtitle, style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_before = Pt(16)
    source_run = source.add_run(f"对应视频：{guide.video_file_name}")
    base.set_run_font(source_run, size=10.5, color=base.MUTED_COLOR)

    duration = document.add_paragraph()
    duration.alignment = WD_ALIGN_PARAGRAPH.CENTER
    duration_run = duration.add_run(f"视频时长：{guide.duration_text}  |  软件：机器视觉AI检测系统 V1.0.0.0")
    base.set_run_font(duration_run, size=10.5, color=base.MUTED_COLOR)

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(28)
    date_run = date.add_run(f"整理日期：{DOCUMENT_DATE_TEXT}")
    base.set_run_font(date_run, size=10, color=base.MUTED_COLOR)
    document.add_page_break()


def add_overview(base, document: Document, guide: GuideSpec, decimal_number_id: int) -> None:
    """添加目标、结构、准备项和步骤概览。"""

    document.add_heading("一、文档目标", level=1)
    document.add_paragraph(guide.goal)

    document.add_heading("二、功能结构", level=1)
    structure = document.add_paragraph()
    structure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    structure.paragraph_format.space_before = Pt(6)
    structure.paragraph_format.space_after = Pt(10)
    structure_run = structure.add_run(guide.structure)
    base.set_run_font(structure_run, size=12, color=base.HEADING_DARK_BLUE, bold=True)

    document.add_heading("三、开始前准备", level=1)
    for label, text in guide.preparations:
        base.add_labeled_paragraph(document, label, text)

    document.add_heading("四、操作步骤概览", level=1)
    for step in guide.steps:
        base.add_numbered_item(document, step.title, decimal_number_id)


def add_guide_step(base, document: Document, guide: GuideSpec, number: int, step: GuideStep) -> None:
    """添加一个包含操作说明、截图、提示和完成标志的步骤页。"""

    image_specs = []
    for screenshot in step.screenshots:
        image_path = prepare_screenshot(guide, screenshot)
        image_specs.append((image_path, screenshot.caption, screenshot.alt_text, screenshot.width_inches))

    base.add_step(
        document,
        number,
        step.title,
        step.actions,
        step.completion,
        image_specs,
        callout=step.callout,
        page_break=True,
    )


def build_guide(base, guide: GuideSpec) -> Path:
    """根据视频说明和关键截图生成一份 Word 操作指南。"""

    source_video = VIDEO_DIRECTORY / guide.video_file_name
    if not source_video.exists():
        raise FileNotFoundError(f"源视频不存在：{source_video}")

    OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)
    logo_path = base.prepare_logo()

    document = Document()
    base.configure_page_geometry(document)
    base.configure_document_styles(document)
    base.configure_header_and_footer(document, logo_path)
    document.core_properties.title = guide.title
    document.core_properties.subject = guide.subtitle
    document.core_properties.author = "TDJS-Vision"
    document.core_properties.keywords = guide.keywords
    document.core_properties.comments = f"根据 {guide.video_file_name} 整理"

    decimal_number_id = base.create_numbering_definition(document, "decimal", "%1.")
    check_number_id = base.create_numbering_definition(document, "bullet", "☐", "Segoe UI Symbol")

    add_cover_page(base, document, guide, logo_path)
    add_overview(base, document, guide, decimal_number_id)

    for index, step in enumerate(guide.steps, start=1):
        add_guide_step(base, document, guide, index, step)

    document.add_heading("最终检查清单", level=1)
    for item in guide.checklist:
        base.add_check_item(document, item, check_number_id)
    base.add_callout(document, "归档建议：", "完成配置或培训后，把本指南、源视频、方案文件、模型文件和现场参数一起归档，便于后续复盘和交接。")

    output_path = OUTPUT_DIRECTORY / guide.output_file_name
    document.save(output_path)
    return output_path


def create_guides() -> list[GuideSpec]:
    """创建三份新增视频对应的文档规格。"""

    return [
        GuideSpec(
            video_stem="框架基本功能介绍",
            video_file_name="框架基本功能介绍.mp4",
            title="机器视觉 AI 检测系统基本功能介绍",
            subtitle="软件入口、方案、日志、流程、设置和训练功能说明",
            output_file_name="框架基本功能介绍操作流程.docx",
            duration_text="04:13",
            goal="本指南把录屏中的软件框架讲解整理为可查阅的功能说明，帮助使用者快速知道常用按钮、菜单、日志、流程编辑、通信设置、检测项和训练入口分别做什么。",
            structure="方案文件 → 主界面与图像窗口 → 程序运行日志 → 流程编辑与节点库\n设备通信 → 检测项管理 → AI配置/画布/系统/信号设置 → AI训练与语言切换",
            preparations=[
                ("软件：", "机器视觉 AI 检测系统已安装并能正常启动。"),
                ("权限：", "需要修改方案、设备或检测项时，建议先登录对应权限账号。"),
                ("方案：", "如果只是学习界面，可使用演示方案；如果修改现场方案，先备份原方案。"),
                ("习惯：", "每次改完关键参数都要保存方案，否则重启后可能丢失。"),
            ],
            steps=[
                GuideStep(
                    title="启动软件并认识主界面",
                    actions=[
                        ("打开：", "启动“机器视觉AI检测系统”，确认标题栏显示当前版本。"),
                        ("识别区域：", "主界面上方是菜单和快捷按钮，中间是检测图像窗口，右侧可显示日志、检测结果和品牌区。"),
                        ("观察状态：", "图像窗口左下角会显示图像尺寸、RGB 或灰度信息、鼠标坐标和缩放比例。"),
                    ],
                    completion="能说清楚菜单区、图像窗口、运行控制区和右侧日志区的位置。",
                    screenshots=[
                        ScreenshotSpec(5000, "01_主界面功能区.jpg", "图 1：软件主界面包含菜单、快捷按钮、图像窗口和日志区域（视频 00:05）", "机器视觉AI检测系统主界面功能区", 0, 1030)
                    ],
                    callout=("注意：", "视频中的窗口布局是演示布局，实际项目可根据显示器数量和工位需要调整。"),
                ),
                GuideStep(
                    title="新建、打开、保存和另存方案",
                    actions=[
                        ("新建方案：", "用于重置软件内部主要参数，适合从空白状态重新搭建流程。"),
                        ("打开方案：", "选择已有 `.Sol` 方案文件，加载其中的流程、窗口和配置。"),
                        ("保存方案：", "把当前软件参数写入方案文件；修改流程、设备、检测项后要主动保存。"),
                        ("另存方案：", "将当前方案复制到新位置，适合做版本备份或派生新项目。"),
                    ],
                    completion="能区分新建、打开、保存、另存的用途，并知道改完参数后必须保存。",
                    screenshots=[
                        ScreenshotSpec(35000, "02_保存方案说明.jpg", "图 2：保存方案用于保存当前软件参数（视频 00:35）", "保存方案功能说明", 0, 1030),
                        ScreenshotSpec(60000, "03_另存方案说明.jpg", "图 3：另存方案可把当前方案保存到新位置（视频 01:00）", "另存方案功能说明", 0, 1030, 5.8),
                    ],
                    callout=("建议：", "现场调试前先另存一份备份方案，再在副本上修改，避免覆盖可运行版本。"),
                ),
                GuideStep(
                    title="查看程序运行日志",
                    actions=[
                        ("打开：", "点击右侧“程序运行日志”区域或日志按钮，查看节点加载、运行、警告和异常记录。"),
                        ("筛选：", "底部可按“全部、消息、调试、警告、异常、致命”分类查看。"),
                        ("定位：", "节点运行失败时优先看致命和异常，再回到对应节点处理参数。"),
                    ],
                    completion="能够通过日志判断节点是否加载成功、运行成功，以及异常发生在哪个节点。",
                    screenshots=[
                        ScreenshotSpec(55000, "04_程序运行日志.jpg", "图 4：程序运行日志记录节点加载、运行和错误信息（视频 00:55）", "程序运行日志与分类筛选", 0, 1030)
                    ],
                    callout=("排查顺序：", "先看红色致命错误，再看黄色警告；不要只看图像窗口是否有画面。"),
                ),
                GuideStep(
                    title="进入流程编辑并认识节点库",
                    actions=[
                        ("进入：", "打开流程编辑界面，顶部标签显示当前流程，旁边的“+ / -”用于增加或删除流程。"),
                        ("节点库：", "左侧按模块分类显示节点，例如图像采集、图像处理、AI检测识别、测量工具、图形创建、设备通信、逻辑工具和结果处理。"),
                        ("运行结果：", "流程节点会显示成功、NG、耗时等状态，右侧小窗口可预览输出图像。"),
                    ],
                    completion="能够找到流程标签、节点库分类、运行按钮、节点状态和输出预览。",
                    screenshots=[
                        ScreenshotSpec(70000, "05_流程编辑节点库.jpg", "图 5：流程编辑界面左侧是算法和工具节点集合（视频 01:10）", "流程编辑界面和节点库分类", 0, 1030)
                    ],
                    callout=("效率：", "节点耗时是调试节拍的重要依据，运行慢时优先查看耗时最高的节点。"),
                ),
                GuideStep(
                    title="管理多个检测内容",
                    actions=[
                        ("打开：", "进入检测项管理，选择检测项文件。"),
                        ("维护：", "可刷新列表、添加、删除选中项、导出和保存。"),
                        ("识别：", "列表中每个检测项代表一个可被 AI 节点或流程引用的检测内容。"),
                    ],
                    completion="检测项列表能显示当前项目需要的检测内容，新增或修改后完成保存。",
                    screenshots=[
                        ScreenshotSpec(185000, "06_检测项管理.jpg", "图 6：检测项管理中可维护多个不同检测内容（视频 03:05）", "检测项管理窗口", 0, 1030)
                    ],
                    callout=("注意：", "检测项名称要和工位、相机、模型保持一致，否则流程可运行但结果可能对应错对象。"),
                ),
                GuideStep(
                    title="使用设备通信和光源相关入口",
                    actions=[
                        ("通信入口：", "工具栏提供 PLC、Modbus、TCP 等设备通信入口。"),
                        ("Modbus：", "进入 Modbus 设备管理后，可添加设备并维护通信参数。"),
                        ("光源管理：", "视频说明光源管理可通过 TCP 或串口控制光源亮度，前提是光源控制器支持对应通信。"),
                    ],
                    completion="知道设备通信入口在哪里，并能区分通信设备管理和光源亮度控制。",
                    screenshots=[
                        ScreenshotSpec(125000, "07_Modbus设备管理.jpg", "图 7：Modbus 设备管理用于维护通信设备和参数（视频 02:05）", "Modbus设备管理窗口", 0, 1030)
                    ],
                    callout=("现场：", "通信类配置会影响外部 PLC、光源或 IO 卡，修改前要确认设备地址和协议。"),
                ),
                GuideStep(
                    title="调整图像窗口画布布局",
                    actions=[
                        ("进入：", "在设置中打开画布设置或图像显示窗口布局。"),
                        ("选择：", "按现场显示需求选择 1、2、3、4 或更多画布组合。"),
                        ("验证：", "布局调整后运行流程，确认每个图像显示节点输出到正确窗口。"),
                    ],
                    completion="图像窗口数量和布局满足现场观察需求，图像显示节点没有输出错窗口。",
                    screenshots=[
                        ScreenshotSpec(195000, "08_图像窗口布局.jpg", "图 8：图像显示窗口布局可选择不同画布数量（视频 03:15）", "图像显示窗口布局设置", 0, 1030)
                    ],
                    callout=("提醒：", "复制流程后常会忘记改图像窗口编号，布局调整后要逐个节点检查输出去向。"),
                ),
                GuideStep(
                    title="配置系统启动和全局信号",
                    actions=[
                        ("系统设置：", "可设置软件自启动、默认打开方案、方案加载完成后自动运行等选项。"),
                        ("全局信号：", "全局信号设置用于维护视觉准备、视觉离线、视觉监听等长期有效的通信信号。"),
                        ("启用：", "每条信号都要选择通信设备、地址、数据类型和数据值，再勾选启用。"),
                    ],
                    completion="系统启动策略和全局信号配置符合现场联机要求。",
                    screenshots=[
                        ScreenshotSpec(210000, "09_系统设置.jpg", "图 9：系统设置可配置自启动、默认方案和自动运行（视频 03:30）", "系统设置窗口", 0, 1030, 5.8),
                        ScreenshotSpec(225000, "10_全局信号设置.jpg", "图 10：全局信号设置包含通信设备、地址、数据类型和启用项（视频 03:45）", "全局信号设置窗口", 0, 1030, 5.8),
                    ],
                    callout=("谨慎：", "全局信号通常与产线联锁相关，改动后需要和 PLC 或上位机一起验证。"),
                ),
                GuideStep(
                    title="进入 AI 训练并切换语言",
                    actions=[
                        ("AI训练：", "菜单中的“AI训练”可进入无监督 AI 训练界面，用于选择训练图像路径、OK/NG 图像和 ROI。"),
                        ("语言：", "语言下拉框可在中文和 English 之间切换。"),
                        ("确认：", "切换语言后检查菜单、按钮和登录入口显示是否符合使用者习惯。"),
                    ],
                    completion="能够打开训练界面，并知道语言切换会影响主界面显示文本。",
                    screenshots=[
                        ScreenshotSpec(240000, "11_AI训练界面.jpg", "图 11：无监督 AI 训练界面包含路径、OK/NG 数量和 ROI 区域（视频 04:00）", "无监督AI训练界面", 0, 1030, 5.8),
                        ScreenshotSpec(250000, "12_英文界面.jpg", "图 12：语言切换后主界面显示为英文（视频 04:10）", "英文语言界面", 0, 1030, 5.8),
                    ],
                    callout=("交接：", "培训文档建议统一使用中文界面截图，现场人员更容易对照操作。"),
                ),
            ],
            checklist=[
                "已确认新建、打开、保存和另存方案的区别。",
                "已知道日志区和不同日志级别的用途。",
                "已能进入流程编辑并识别常用节点分类。",
                "已能找到检测项管理并理解检测项与工位的关系。",
                "已了解 PLC、Modbus、TCP、光源管理等通信入口。",
                "已能调整图像窗口画布布局。",
                "已了解系统自启动、默认方案、自动运行和全局信号配置。",
                "已知道 AI 训练入口和语言切换方式。",
            ],
        ),
        GuideSpec(
            video_stem="模型加密与获取视频",
            video_file_name="模型加密与获取视频.mp4",
            title="模型获取与模型加密操作流程",
            subtitle="源模型定位、ONNX 转 Engine、模型加密与交付文件说明",
            output_file_name="模型加密与获取操作流程.docx",
            duration_text="03:26",
            goal="本指南把录屏中的模型获取、GPU Engine 转换、ONNX/Engine 加密和交付文件要求整理成可执行步骤，避免上线时漏传模型或传错模型类型。",
            structure="网络模型目录 → 选择目标模型 → 区分 CPU/GPU 文件\nONNX 转 TensorRT Engine → 生成加密文件 → 按运行设备整理交付文件",
            preparations=[
                ("模型来源：", "能访问共享目录或项目模型目录，例如 `public` 网络盘中的模型文件夹。"),
                ("工具：", "准备 ONNX 转 TensorRT Engine 工具和 ONNX 模型加密工具。"),
                ("路径：", "建议把待处理模型复制到本机简单路径，例如桌面的 `Model` 文件夹。"),
                ("硬件：", "确认目标设备使用 CPU 还是 GPU；GPU 模型还要确认显卡型号和 TensorRT 环境。"),
            ],
            steps=[
                GuideStep(
                    title="进入共享模型目录并选择目标模型",
                    actions=[
                        ("打开：", "进入网络共享或项目模型目录，找到与当前项目、类别和日期匹配的模型文件夹。"),
                        ("选择：", "视频中进入端子模型的 12 类模型目录，并选择最新日期的模型版本。"),
                        ("核对：", "优先确认模型类别、训练日期、客户项目和工位需求一致。"),
                    ],
                    completion="已经定位到本次需要处理的模型目录，不再从旧版本或错误类别中取文件。",
                    screenshots=[
                        ScreenshotSpec(15000, "01_选择模型目录.jpg", "图 1：在共享模型目录中选择目标模型版本（视频 00:15）", "共享目录中的端子模型和12类模型文件夹", 0, 1030)
                    ],
                    callout=("建议：", "模型目录命名通常带日期或客户信息，复制前先和项目记录核对，避免同名类别拿错版本。"),
                ),
                GuideStep(
                    title="区分 CPU 模型和 GPU 模型",
                    actions=[
                        ("CPU：", "视频说明 CPU 模型使用 ONNX 文件，常见文件为 `*.onnx` 或加密后的 `*.onnx.encrypted`。"),
                        ("GPU：", "GPU 模型需要 Engine 文件，通常还要根据显卡型号选择对应目录，例如 1050ti、3050 等。"),
                        ("源文件：", "为后续转换和加密，先找一个源 ONNX 模型文件作为处理对象。"),
                    ],
                    completion="明确当前项目要交付 CPU 文件、GPU 文件，还是两类文件都要交付。",
                    screenshots=[
                        ScreenshotSpec(25000, "02_CPU和GPU模型文件.jpg", "图 2：模型目录中包含不同显卡目录和 ONNX 文件（视频 00:25）", "模型目录中的1050ti、3050和ONNX文件", 0, 1030)
                    ],
                    callout=("规则：", "CPU 场景通常只需要加密 ONNX；GPU 场景通常需要 ONNX 加密文件和 Engine 加密文件同时交付。"),
                ),
                GuideStep(
                    title="准备本机模型工作目录",
                    actions=[
                        ("复制：", "把选中的源 ONNX 模型复制到本机工作目录，视频中使用桌面 `Model` 文件夹。"),
                        ("命名：", "保持文件名简单，例如 `best.onnx`，便于转换工具和加密工具读取。"),
                        ("检查：", "确认文件大小正常，路径中不要出现临时下载未完成或权限受限的位置。"),
                    ],
                    completion="本机 `Model` 目录中已有待转换或待加密的 ONNX 模型。",
                    screenshots=[
                        ScreenshotSpec(95000, "03_本机模型目录.jpg", "图 3：把源模型准备到本机 Model 目录（视频 01:35）", "本机Model目录中的ONNX模型", 0, 1030)
                    ],
                    callout=("效率：", "模型转换和加密建议在本机磁盘执行，网络盘读写慢或断连会导致转换失败。"),
                ),
                GuideStep(
                    title="打开 ONNX 转 TensorRT Engine 工具",
                    actions=[
                        ("打开工具：", "启动“ONNX 转 TensorRT Engine”工具。"),
                        ("选择输入：", "点击 ONNX 文件后的“浏览”，选择本机 `best.onnx`。"),
                        ("选择输出：", "设置 Engine 输出路径，例如 `best.engine`。"),
                        ("参数：", "Workspace 视频示例为 2048 MiB；默认关闭 FP16 和 TF32，优先保证结果复现。"),
                    ],
                    completion="工具中已填写 ONNX 输入路径、Engine 输出路径和转换参数。",
                    screenshots=[
                        ScreenshotSpec(135000, "04_ONNX转Engine工具.jpg", "图 4：ONNX 转 TensorRT Engine 工具参数区（视频 02:15）", "ONNX转TensorRT Engine工具窗口", 0, 1030)
                    ],
                    callout=("性能：", "启用 FP16 可提升速度但可能影响精度；没有确认前按视频默认方式处理更稳。"),
                ),
                GuideStep(
                    title="开始转换并等待 Engine 生成",
                    actions=[
                        ("启动：", "点击“开始转换”。"),
                        ("等待：", "转换过程可能较久，日志会显示解析 ONNX、构建 TensorRT engine、Engine bytes 等信息。"),
                        ("完成：", "出现“Engine 转换完成”提示后，确认输出目录中生成 `best.engine`。"),
                    ],
                    completion="本机模型目录中已有转换完成的 `best.engine` 文件。",
                    screenshots=[
                        ScreenshotSpec(155000, "05_Engine转换完成.jpg", "图 5：日志提示 Engine 转换完成并生成 best.engine（视频 02:35）", "ONNX转Engine工具显示转换完成", 0, 1030)
                    ],
                    callout=("不要中断：", "转换期间不要关闭工具或移动源模型；若失败，先保留日志再检查模型、显卡和 TensorRT 环境。"),
                ),
                GuideStep(
                    title="使用模型加密工具生成加密文件",
                    actions=[
                        ("打开工具：", "启动 ONNX 模型加密工具。"),
                        ("选择输入：", "选择需要加密的源文件；CPU 使用 ONNX，GPU 也要对 Engine 加密。"),
                        ("生成：", "点击“生成加密文件”，等待提示“加密完成”。"),
                        ("输出：", "工具会生成对应的 `.encrypted` 文件，例如 `best.engine.encrypted`。"),
                    ],
                    completion="需要加密的模型均已生成 `.encrypted` 文件。",
                    screenshots=[
                        ScreenshotSpec(190000, "06_模型加密完成.jpg", "图 6：模型加密工具提示加密完成（视频 03:10）", "ONNX模型加密工具加密完成", 0, 1030)
                    ],
                    callout=("命名：", "加密输出文件建议保留原模型名，只追加 `.encrypted`，便于后续配置文件引用。"),
                ),
                GuideStep(
                    title="按运行设备整理交付文件",
                    actions=[
                        ("CPU 设备：", "上传或交付 `best.onnx.encrypted`。"),
                        ("GPU 设备：", "上传或交付 `best.onnx.encrypted` 与 `best.engine.encrypted` 两个文件。"),
                        ("核对：", "源 `best.onnx` 和未加密 `best.engine` 可留作开发归档，现场运行优先使用加密文件。"),
                    ],
                    completion="交付目录中包含目标设备实际需要的加密模型文件，且文件名清晰可追溯。",
                    screenshots=[
                        ScreenshotSpec(200000, "07_交付文件说明.jpg", "图 7：CPU 和 GPU 模型需要上传的文件不同（视频 03:20）", "Model目录中的best.engine、best.engine.encrypted和best.onnx.encrypted", 0, 1030)
                    ],
                    callout=("上线前：", "把加密模型路径同步到 AI 配置文件中的 ModelPath，再在软件日志中确认模型加载成功。"),
                ),
            ],
            checklist=[
                "已确认模型目录、客户项目、类别和日期版本正确。",
                "已区分 CPU 模型和 GPU 模型交付内容。",
                "已把源 ONNX 复制到本机稳定目录。",
                "已完成 ONNX 到 Engine 转换，并确认生成 `best.engine`。",
                "已生成 `best.onnx.encrypted`。",
                "GPU 场景已生成 `best.engine.encrypted`。",
                "已按 CPU/GPU 设备要求整理交付文件。",
                "已准备后续在 AI 配置文件中引用加密模型路径。",
            ],
            keywords="机器视觉,模型加密,ONNX,TensorRT,Engine,AI检测",
        ),
        GuideSpec(
            video_stem="科瑞IO卡前期连接工作",
            video_file_name="科瑞IO卡前期连接工作.mp4",
            title="科瑞 IO 卡前期连接操作流程",
            subtitle="网络连通性检查、网卡 IP 配置和软件启动验证",
            output_file_name="科瑞IO卡前期连接操作流程.docx",
            duration_text="01:35",
            goal="本指南把录屏中的科瑞 IO 卡前期通信检查整理为标准步骤，用于在软件接入 IO 卡前先确认电脑网卡和 IO 卡处在同一网段，并能稳定 ping 通默认地址。",
            structure="确认默认地址 → ping 连通性测试 → 打开网络连接\n选择正确网卡 → 配置 IPv4 同网段地址 → 再次 ping 验证 → 启动视觉软件",
            preparations=[
                ("默认地址：", "视频中科瑞 IO 卡默认地址为 `192.168.123.199`。"),
                ("网线：", "IO 卡和电脑网卡已连接，网口指示灯正常。"),
                ("权限：", "修改 Windows 网卡 IP 需要管理员权限或现场允许的维护权限。"),
                ("原则：", "电脑 IP 要和 IO 卡在同一网段，但不能与 IO 卡地址重复。"),
            ],
            steps=[
                GuideStep(
                    title="用 ping 检查默认地址",
                    actions=[
                        ("打开：", "进入 Windows 命令提示符。"),
                        ("输入：", "执行 `ping 192.168.123.199`，检查 IO 卡默认地址是否连通。"),
                        ("观察：", "若出现“请求超时”，说明当前电脑与 IO 卡暂时不通，需要检查网卡网段或连接。"),
                    ],
                    completion="已经获得第一次 ping 结果，并能判断当前是否连通。",
                    screenshots=[
                        ScreenshotSpec(25000, "01_ping请求超时.jpg", "图 1：ping 默认地址超时，说明当前连接或网段有问题（视频 00:25）", "命令提示符中ping 192.168.123.199请求超时", 0, 930)
                    ],
                    callout=("先判断：", "ping 不通不一定是 IO 卡坏了，视频中优先排查电脑和 IO 卡是否在同一网段。"),
                ),
                GuideStep(
                    title="打开网络连接并选择正确网卡",
                    actions=[
                        ("进入：", "打开“网络和 Internet → 网络连接”。"),
                        ("选择：", "在多个以太网适配器中找到实际连接 IO 卡的网卡。"),
                        ("判断：", "可通过网线状态、连接速度和网口名称判断；视频中特别提示看速度，正常应为千兆连接。"),
                    ],
                    completion="已选中连接 IO 卡的网卡，没有误改无线网卡或其他产线网卡。",
                    screenshots=[
                        ScreenshotSpec(40000, "02_网络连接列表.jpg", "图 2：在网络连接中选择连接 IO 卡的以太网适配器（视频 00:40）", "Windows网络连接列表和以太网适配器", 0, 930)
                    ],
                    callout=("注意：", "电脑可能有多块网卡，只改连接 IO 卡的那一块；误改其他网卡会影响相机或外网通信。"),
                ),
                GuideStep(
                    title="进入 IPv4 属性",
                    actions=[
                        ("属性：", "右键目标以太网，打开“属性”。"),
                        ("协议：", "选中“Internet 协议版本 4 (TCP/IPv4)”，点击“属性”。"),
                        ("切换：", "从自动获取 IP 改为“使用下面的 IP 地址”。"),
                    ],
                    completion="已打开目标网卡的 IPv4 属性窗口，准备填写静态 IP。",
                    screenshots=[
                        ScreenshotSpec(50000, "03_IPv4属性.jpg", "图 3：进入 TCP/IPv4 属性并切换为手动 IP（视频 00:50）", "Windows TCP IPv4属性窗口", 0, 930)
                    ],
                    callout=("范围：", "只改 IPv4 地址和子网掩码即可；DNS 和默认网关通常不是本次 IO 卡直连的重点。"),
                ),
                GuideStep(
                    title="配置与 IO 卡同网段的电脑 IP",
                    actions=[
                        ("填写：", "电脑 IP 填同一网段但不重复的地址，例如 `192.168.123.1`。"),
                        ("掩码：", "子网掩码通常填写 `255.255.255.0`。"),
                        ("保存：", "点击确定保存所有属性窗口，等待网络适配器应用新地址。"),
                    ],
                    completion="电脑网卡已经改为与 `192.168.123.199` 同网段的静态地址。",
                    screenshots=[
                        ScreenshotSpec(60000, "04_配置静态IP.jpg", "图 4：把电脑网卡设置到 IO 卡同一网段（视频 01:00）", "IPv4属性窗口中填写静态IP地址", 0, 930)
                    ],
                    callout=("不要重复：", "电脑 IP 不能也填 `192.168.123.199`，否则会和 IO 卡冲突。"),
                ),
                GuideStep(
                    title="再次 ping 验证通信",
                    actions=[
                        ("重试：", "回到命令提示符，再次执行 `ping 192.168.123.199`。"),
                        ("成功标志：", "出现“来自 192.168.123.199 的回复”，并显示字节、时间和 TTL。"),
                        ("失败处理：", "仍超时时继续检查网线、目标网卡、IP 段、IO 卡电源和默认地址是否被改过。"),
                    ],
                    completion="ping 已收到回复，说明电脑和科瑞 IO 卡网络通信已通。",
                    screenshots=[
                        ScreenshotSpec(70000, "05_ping通信成功.jpg", "图 5：出现来自默认地址的回复，说明通信已通（视频 01:10）", "命令提示符中ping 192.168.123.199返回TTL", 0, 930)
                    ],
                    callout=("验收：", "连续 ping 多次都稳定回复后，再进入视觉软件配置 IO 通信。"),
                ),
                GuideStep(
                    title="启动视觉软件并观察异常提示",
                    actions=[
                        ("启动：", "打开机器视觉 AI 检测系统，等待方案加载完成。"),
                        ("观察：", "若弹出“相机对象为空”等提示，说明软件方案中相机或节点配置还未完成，和 IO 卡网络连通不是同一个问题。"),
                        ("后续：", "网络已通后，再到软件里的 IO 或通信节点配置科瑞卡地址、端口和信号。"),
                    ],
                    completion="软件已启动，且能把网络连通问题与软件内部相机/节点配置问题分开处理。",
                    screenshots=[
                        ScreenshotSpec(85000, "06_软件启动提示.jpg", "图 6：网络验证后启动视觉软件并观察方案加载提示（视频 01:25）", "机器视觉AI检测系统启动并弹出相机对象为空提示", 0, 1000)
                    ],
                    callout=("边界：", "本视频只完成前期网络连接检查，软件内 IO 节点参数仍需在项目流程中另行配置。"),
                ),
            ],
            checklist=[
                "已确认科瑞 IO 卡默认地址为 `192.168.123.199`。",
                "已用 ping 做首次连通性测试。",
                "已选择实际连接 IO 卡的以太网适配器。",
                "电脑静态 IP 已设置到 `192.168.123.x` 网段，且未与 IO 卡地址冲突。",
                "子网掩码已设置为 `255.255.255.0`。",
                "再次 ping 已出现 `来自 192.168.123.199 的回复`。",
                "已确认软件启动提示和 IO 卡网络连接属于不同排查方向。",
            ],
            keywords="机器视觉,科瑞IO卡,网络连接,ping,IPv4",
        ),
    ]


def main() -> None:
    """执行三份新增视频文档的生成流程。"""

    base = load_base_helpers()
    output_paths = [build_guide(base, guide) for guide in create_guides()]
    for output_path in output_paths:
        print(output_path)


if __name__ == "__main__":
    main()
