"""把真实流程截图、精确红框和逐项说明装配为独立 Word 说明书。"""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

from build_real_workflow_annotations import ModuleAnnotation, load_annotations
from real_workflow_capture_model import load_module_review


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_SOURCE = ROOT / "outputs" / "TDJS-Vision工业视觉应用框架功能与应用说明书-图文增强版.docx"
DEFAULT_OUTPUT = ROOT / "outputs" / "TDJS-Vision工业视觉应用框架功能与应用说明书-真实流程演示版.docx"
DEFAULT_CONFIG = ROOT / "docs" / "manual-assets" / "real-workflow" / "annotations.json"
DEFAULT_REVIEW = ROOT / "docs" / "manual-assets" / "real-workflow" / "module-review.csv"
DEFAULT_ANNOTATED = ROOT / "docs" / "manual-assets" / "real-workflow" / "annotated-screenshots"
CIRCLED_NUMBERS = {1: "①", 2: "②", 3: "③", 4: "④"}
HARDWARE_NOTE = "配置演示，当前未连接真实设备；现场使用时按实际品牌、地址、端口和站号完成配置。"
THREE_D_NOTE = "当前三维能力边界：三维相机取图、3D图像源和3D图像显示；暂不宣称三维测量或三维缺陷算法能力。"


@dataclass(frozen=True)
class BuildReport:
    """记录 Word 一一替换与源文件保护结果。"""

    modules_replaced: int
    captions_written: int
    legends_written: int
    source_unchanged: bool
    errors: tuple[str, ...]


def _sha256(path: Path) -> str:
    """计算文件哈希，用于确认源 Word 未被修改。"""

    return hashlib.sha256(path.read_bytes()).hexdigest()


def _module_name(heading_text: str) -> str:
    """从带章节号的二级标题提取模块名称。"""

    return re.sub(r"^\s*\d+(?:\.\d+)*\s*[　 ]+\s*", "", heading_text).strip()


def _safe_name(module_name: str) -> str:
    """生成标注图对应的安全文件名。"""

    return re.sub(r'[<>:"/\\|?*]', "_", module_name).strip() or "未命名模块"


def _image_paths(config: dict[str, ModuleAnnotation]) -> dict[str, Path]:
    """按配置顺序定位全部标注图。"""

    return {
        name: DEFAULT_ANNOTATED / f"{index:03d}-{_safe_name(name)}.png"
        for index, name in enumerate(config, start=1)
    }


def _legend(annotation: ModuleAnnotation) -> str:
    """生成红框编号与使用对象的一一对应说明。"""

    return "标注说明：" + "　".join(
        f"{CIRCLED_NUMBERS[region.index]} {region.label}"
        for region in annotation.regions
    )


def _write_caption(
    paragraph: Paragraph,
    module_name: str,
    annotation: ModuleAnnotation,
    capture_kind: str,
) -> None:
    """写入统一图注、编号说明和真实能力边界。"""

    paragraph.text = f"图：{module_name}真实演示界面"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    legend_run = paragraph.add_run("\n" + _legend(annotation))
    legend_run.bold = True
    legend_run.font.size = Pt(9)
    legend_run.font.color.rgb = RGBColor(198, 0, 18)
    if capture_kind == "硬件配置":
        note = HARDWARE_NOTE
    elif module_name in {"3D图像源", "3D图像显示", "三维相机"}:
        note = THREE_D_NOTE
    elif module_name == "运行参数":
        note = "运行参数表按流程中开放动态调节的节点自动生成；当前演示流程展示入口、字段结构和保存操作。"
    else:
        note = ""
    if note:
        note_run = paragraph.add_run("\n说明：" + note)
        note_run.italic = True
        note_run.font.size = Pt(9)


def _insert_figure_before(
    document: Document,
    anchor: Paragraph,
    module_name: str,
    annotation: ModuleAnnotation,
    image_path: Path,
    capture_kind: str,
) -> None:
    """在原章节缺图时，于使用说明前补入真实截图及图注。"""

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.3))
    caption = document.add_paragraph()
    _write_caption(caption, module_name, annotation, capture_kind)
    anchor._p.addprevious(image_paragraph._p)
    anchor._p.addprevious(caption._p)


def _ensure_alt_text(document: Document, modules: set[str]) -> None:
    """为每张说明书图片补充与所在模块一致的中文替代文本。"""

    current_module: str | None = None
    for child in document.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        if paragraph.style.name == "Heading 2":
            current_module = _module_name(paragraph.text)
        description = (
            f"{current_module}真实界面红框标注图"
            if current_module in modules
            else "TDJS-Vision工业视觉应用框架图片"
        )
        for properties in paragraph._p.xpath(".//wp:docPr"):
            properties.set("descr", description)


def _mark_text_approved(review_path: Path) -> None:
    """装配成功后更新模块核对清单的文字状态。"""

    with review_path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    fieldnames = list(rows[0])
    for row in rows:
        row["文字合格"] = "是"
    with review_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def build_real_workflow_manual(
    source_docx: Path,
    output_docx: Path,
    annotations_path: Path,
    review_path: Path,
) -> BuildReport:
    """使用 86 张已批准标注图生成独立的真实流程演示版说明书。"""

    source_hash = _sha256(source_docx)
    config = load_annotations(annotations_path)
    review = load_module_review(review_path)
    if set(config) != set(review):
        raise ValueError("标注配置与模块核对清单范围不一致。")
    for name, row in review.items():
        if not (row.screenshot_approved and row.annotation_approved):
            raise ValueError(f"模块“{name}”的截图或红框尚未批准。")
    images = _image_paths(config)
    missing_images = [name for name, path in images.items() if not path.exists()]
    if missing_images:
        raise FileNotFoundError("缺少模块标注图：" + "、".join(missing_images))

    document = Document(source_docx)
    processed: set[str] = set()
    current_module: str | None = None
    pending_caption: str | None = None
    captions_written = 0
    legends_written = 0

    for child in list(document.element.body.iterchildren()):
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        if paragraph.style.name == "Heading 2":
            current_module = _module_name(paragraph.text)
            pending_caption = None
            continue

        if pending_caption and paragraph.text.strip().startswith(("图：", "图:")):
            name = pending_caption
            _write_caption(paragraph, name, config[name], review[name].capture_kind)
            captions_written += 1
            legends_written += 1
            pending_caption = None
            continue

        if (
            paragraph.style.name == "Heading 3"
            and current_module in config
            and current_module not in processed
        ):
            _insert_figure_before(
                document,
                paragraph,
                current_module,
                config[current_module],
                images[current_module],
                review[current_module].capture_kind,
            )
            processed.add(current_module)
            captions_written += 1
            legends_written += 1
            pending_caption = None
            continue

        if current_module not in config or current_module in processed:
            continue
        blips = paragraph._p.xpath(".//a:blip")
        if not blips:
            continue
        relationship_id, _ = document.part.get_or_add_image(str(images[current_module]))
        blips[0].set(qn("r:embed"), relationship_id)
        processed.add(current_module)
        pending_caption = current_module

    missing_modules = sorted(set(config) - processed)
    errors = tuple(f"模块未替换：{name}" for name in missing_modules)
    _ensure_alt_text(document, set(config))
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    source_unchanged = source_hash == _sha256(source_docx)
    if not errors and captions_written == 86 and legends_written == 86:
        _mark_text_approved(review_path)
    return BuildReport(
        modules_replaced=len(processed),
        captions_written=captions_written,
        legends_written=legends_written,
        source_unchanged=source_unchanged,
        errors=errors,
    )


def main() -> None:
    """使用默认路径生成最终真实流程演示版 Word。"""

    report = build_real_workflow_manual(
        DEFAULT_SOURCE, DEFAULT_OUTPUT, DEFAULT_CONFIG, DEFAULT_REVIEW
    )
    print(
        f"MODULES_REPLACED={report.modules_replaced} "
        f"CAPTIONS={report.captions_written} LEGENDS={report.legends_written} "
        f"ERRORS={len(report.errors)} SOURCE_UNCHANGED={str(report.source_unchanged).lower()}"
    )
    for error in report.errors:
        print("ERROR=" + error)


if __name__ == "__main__":
    main()
