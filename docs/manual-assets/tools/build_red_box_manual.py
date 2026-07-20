"""将模块红框标注图片和编号说明写入 Word 副本。"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Inches, Pt, RGBColor

from build_annotated_screenshots import safe_file_name
from red_box_annotation import ModuleAnnotation, load_annotation_config


ROOT = Path(__file__).resolve().parents[3]
DEFAULT_SOURCE = ROOT / "outputs" / "TDJS-Vision工业视觉应用框架功能与应用说明书-图文增强版.docx"
DEFAULT_OUTPUT = ROOT / "outputs" / "TDJS-Vision工业视觉应用框架功能与应用说明书-红框标注版.docx"
DEFAULT_CONFIG = ROOT / "docs" / "manual-assets" / "annotations" / "red-box-annotations.json"
DEFAULT_ANNOTATED_DIR = ROOT / "docs" / "manual-assets" / "annotated-screenshots"
CIRCLED_NUMBERS = {1: "①", 2: "②", 3: "③", 4: "④"}
LEGEND_PREFIX = "标注说明："


@dataclass(frozen=True)
class BuildReport:
    """记录红框标注版 Word 的装配结果。"""

    annotated_modules: int
    missing_modules: tuple[str, ...]
    source_unchanged: bool


def sha256(path: Path) -> str:
    """计算文件 SHA-256。"""

    return hashlib.sha256(path.read_bytes()).hexdigest()


def extract_module_name(heading_text: str) -> str:
    """从带章节号的二级标题中提取模块名称。"""

    return re.sub(r"^\s*\d+(?:\.\d+)*\s*[　 ]+\s*", "", heading_text).strip()


def annotated_image_paths(
    config: dict[str, ModuleAnnotation],
    annotated_dir: Path,
) -> dict[str, Path]:
    """按配置顺序解析每个模块的标注图片路径。"""

    return {
        module_name: annotated_dir / f"{sequence:03d}-{safe_file_name(module_name)}.png"
        for sequence, module_name in enumerate(config, start=1)
    }


def legend_text(annotation: ModuleAnnotation) -> str:
    """生成红框编号对应的图下注释。"""

    items = [
        f"{CIRCLED_NUMBERS[region.index]} {region.label}"
        for region in annotation.regions
    ]
    return LEGEND_PREFIX + "　".join(items)


def append_or_replace_legend(paragraph: Paragraph, annotation: ModuleAnnotation) -> None:
    """在原图注后追加或替换编号说明行。"""

    base_text = paragraph.text.split(f"\n{LEGEND_PREFIX}", 1)[0]
    if paragraph.text != base_text:
        paragraph.text = base_text
    run = paragraph.add_run("\n" + legend_text(annotation))
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(198, 0, 18)


def insert_annotated_figure_before(
    document: Document,
    anchor: Paragraph,
    module_name: str,
    annotation: ModuleAnnotation,
    image_path: Path,
) -> None:
    """在指定段落前补入缺失的模块截图、图注和编号说明。"""

    image_paragraph = document.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(6.3))

    caption = document.add_paragraph(
        f"图：{module_name}实际界面（由 Release 程序运行或真实窗体加载获得）"
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    append_or_replace_legend(caption, annotation)

    anchor._p.addprevious(image_paragraph._p)
    anchor._p.addprevious(caption._p)


def ensure_image_alt_text(
    document: Document,
    config: dict[str, ModuleAnnotation],
) -> None:
    """为目标文档中缺少说明的图片补充中文替代文本。"""

    current_module: str | None = None
    for child in document.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        if paragraph.style.name == "Heading 2":
            current_module = extract_module_name(paragraph.text)

        description = (
            f"{current_module}界面红框标注图"
            if current_module in config
            else "TDJS-Vision工业视觉应用框架图片"
        )
        for document_properties in paragraph._p.xpath(".//wp:docPr"):
            if not (
                document_properties.get("descr", "").strip()
                or document_properties.get("title", "").strip()
            ):
                document_properties.set("descr", description)


def build_red_box_manual(
    source_docx: Path,
    output_docx: Path,
    config_path: Path,
    annotated_dir: Path,
) -> BuildReport:
    """基于当前说明书生成独立的红框标注版 Word。"""

    source_hash = sha256(source_docx)
    config = load_annotation_config(config_path)
    images = annotated_image_paths(config, annotated_dir)
    missing_image_files = [
        module_name
        for module_name, image_path in images.items()
        if not image_path.exists()
    ]
    if missing_image_files:
        raise FileNotFoundError("缺少模块标注图：" + "、".join(missing_image_files))

    document = Document(source_docx)
    processed: set[str] = set()
    current_module: str | None = None
    pending_caption: str | None = None

    for child in document.element.body.iterchildren():
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document._body)
        if paragraph.style.name == "Heading 2":
            current_module = extract_module_name(paragraph.text)
            pending_caption = None
            continue

        if (
            paragraph.style.name == "Heading 3"
            and current_module in config
            and current_module not in processed
        ):
            insert_annotated_figure_before(
                document,
                paragraph,
                current_module,
                config[current_module],
                images[current_module],
            )
            processed.add(current_module)
            pending_caption = None
            continue

        if pending_caption and paragraph.text.strip().startswith(("图：", "图:")):
            append_or_replace_legend(paragraph, config[pending_caption])
            pending_caption = None
            continue

        if current_module not in config or current_module in processed:
            continue
        blips = paragraph._p.xpath(".//a:blip")
        if not blips:
            continue

        relationship_id, _ = document.part.get_or_add_image(str(images[current_module]))
        blips[0].set(qn("r:embed"), relationship_id)
        processed.add(current_module)
        pending_caption = current_module

    missing_modules = tuple(sorted(set(config) - processed))
    ensure_image_alt_text(document, config)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    return BuildReport(
        annotated_modules=len(processed),
        missing_modules=missing_modules,
        source_unchanged=source_hash == sha256(source_docx),
    )


def main() -> None:
    """使用工程默认路径生成红框标注版说明书。"""

    report = build_red_box_manual(
        DEFAULT_SOURCE,
        DEFAULT_OUTPUT,
        DEFAULT_CONFIG,
        DEFAULT_ANNOTATED_DIR,
    )
    print(
        f"ANNOTATED_MODULES={report.annotated_modules} "
        f"MISSING={len(report.missing_modules)} "
        f"SOURCE_UNCHANGED={str(report.source_unchanged).lower()}"
    )
    if report.missing_modules:
        print("MISSING_MODULES=" + "、".join(report.missing_modules))


if __name__ == "__main__":
    main()
