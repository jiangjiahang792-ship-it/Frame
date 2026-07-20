"""红框标注版 Word 装配测试。"""

from __future__ import annotations

import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw


TOOLS_DIR = Path(__file__).resolve().parents[1]
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

from build_red_box_manual import build_red_box_manual


def sha256(path: Path) -> str:
    """计算文件 SHA-256。"""

    return hashlib.sha256(path.read_bytes()).hexdigest()


def create_marked_image(path: Path, color: str) -> None:
    """创建用于区分图片关系的测试图。"""

    image = Image.new("RGB", (600, 360), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, 560, 320), outline=color, width=12)
    image.save(path)


def missing_alt_text_count(document: Document) -> int:
    """统计未配置替代说明的内嵌图片数量。"""

    return sum(
        1
        for doc_properties in document.element.body.iter(qn("wp:docPr"))
        if not (
            doc_properties.get("descr", "").strip()
            or doc_properties.get("title", "").strip()
        )
    )


class RedBoxManualBuildTests(unittest.TestCase):
    """验证按模块替换图片和追加图下注释。"""

    def test_inserts_image_and_caption_when_module_has_no_source_screenshot(self) -> None:
        """原章节没有截图时，必须在使用方法前补入标注图和对应说明。"""

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_image = root / "source.png"
            create_marked_image(source_image, "gray")
            source_docx = root / "source.docx"
            output_docx = root / "output.docx"
            document = Document()
            document.add_heading("3.1　方案管理", level=2)
            document.add_paragraph("建立、打开、保存和切换完整视觉项目。")
            document.add_paragraph("")
            document.add_heading("使用方法", level=3)
            document.save(source_docx)
            source_hash = sha256(source_docx)

            annotated_dir = root / "annotated"
            annotated_dir.mkdir()
            create_marked_image(annotated_dir / "001-方案管理.png", "red")
            config_path = root / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "方案管理": {
                            "source_file": source_image.name,
                            "regions": [
                                {
                                    "index": 1,
                                    "label": "新建、打开和保存方案",
                                    "box": [0.1, 0.1, 0.9, 0.9],
                                }
                            ],
                        }
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            report = build_red_box_manual(
                source_docx,
                output_docx,
                config_path,
                annotated_dir,
            )

            self.assertEqual(source_hash, sha256(source_docx))
            self.assertEqual(1, report.annotated_modules)
            self.assertEqual(0, len(report.missing_modules))
            output = Document(output_docx)
            text = "\n".join(paragraph.text for paragraph in output.paragraphs)
            self.assertIn("图：方案管理实际界面", text)
            self.assertIn("标注说明：① 新建、打开和保存方案", text)
            relationship_ids = [
                blip.get(qn("r:embed"))
                for blip in output.element.body.iter(qn("a:blip"))
            ]
            self.assertEqual(1, len(relationship_ids))
            self.assertEqual(0, missing_alt_text_count(output))

    def test_replaces_shared_image_per_module_and_preserves_source(self) -> None:
        """共享原图的两个模块必须获得独立标注和正确说明。"""

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_image = root / "source.png"
            create_marked_image(source_image, "gray")
            source_docx = root / "source.docx"
            output_docx = root / "output.docx"
            document = Document()
            for number, module_name in enumerate(("模块甲", "模块乙"), start=1):
                document.add_heading(f"3.{number}　{module_name}", level=2)
                paragraph = document.add_paragraph()
                paragraph.add_run().add_picture(str(source_image), width=Inches(3.0))
                document.add_paragraph(f"图：{module_name}实际界面")
            document.save(source_docx)
            source_hash = sha256(source_docx)

            annotated_dir = root / "annotated"
            annotated_dir.mkdir()
            create_marked_image(annotated_dir / "001-模块甲.png", "red")
            create_marked_image(annotated_dir / "002-模块乙.png", "blue")
            config_path = root / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "模块甲": {
                            "source_file": source_image.name,
                            "regions": [
                                {"index": 1, "label": "甲输入", "box": [0.1, 0.1, 0.9, 0.9]}
                            ],
                        },
                        "模块乙": {
                            "source_file": source_image.name,
                            "regions": [
                                {"index": 1, "label": "乙执行", "box": [0.1, 0.1, 0.9, 0.9]}
                            ],
                        },
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            report = build_red_box_manual(
                source_docx,
                output_docx,
                config_path,
                annotated_dir,
            )

            self.assertEqual(source_hash, sha256(source_docx))
            self.assertEqual(2, report.annotated_modules)
            self.assertEqual(0, len(report.missing_modules))
            output = Document(output_docx)
            text = "\n".join(paragraph.text for paragraph in output.paragraphs)
            self.assertIn("标注说明：① 甲输入", text)
            self.assertIn("标注说明：① 乙执行", text)
            relationship_ids = [
                blip.get(qn("r:embed"))
                for blip in output.element.body.iter(qn("a:blip"))
            ]
            self.assertEqual(2, len(relationship_ids))
            self.assertNotEqual(relationship_ids[0], relationship_ids[1])
            self.assertEqual(0, missing_alt_text_count(output))


if __name__ == "__main__":
    unittest.main()
