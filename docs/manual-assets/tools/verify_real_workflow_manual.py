"""对真实流程演示版说明书执行最终一致性断言。"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from build_real_workflow_annotations import load_annotations
from real_workflow_capture_model import load_module_review


ROOT = Path(__file__).resolve().parents[3]
REAL_WORKFLOW = ROOT / "docs" / "manual-assets" / "real-workflow"
FINAL_DOCX = ROOT / "outputs" / "TDJS-Vision工业视觉应用框架功能与应用说明书-真实流程演示版.docx"
FINAL_PDF = ROOT / ".codex-tmp" / "TDJS-Vision真实流程演示版.pdf"
RENDERED_PAGES = ROOT / ".codex-tmp" / "tdjs-real-workflow-render"


def main() -> None:
    """核对 86 模块、流程状态、Word 结构和渲染页数。"""

    review = load_module_review(REAL_WORKFLOW / "module-review.csv")
    annotations = load_annotations(REAL_WORKFLOW / "annotations.json")
    results = json.loads((REAL_WORKFLOW / "capture-result.json").read_text(encoding="utf-8"))
    canvas = next(item for item in results if item["module"] == "流程画布")
    annotated_images = list((REAL_WORKFLOW / "annotated-screenshots").glob("[0-9][0-9][0-9]-*.png"))
    document = Document(FINAL_DOCX)
    captions = sum(
        paragraph.text.startswith("图：") and "真实演示界面" in paragraph.text
        for paragraph in document.paragraphs
    )
    legends = sum("标注说明：" in paragraph.text for paragraph in document.paragraphs)
    pdf_pages = len(PdfReader(FINAL_PDF).pages)
    rendered_pages = len(list(RENDERED_PAGES.glob("page-*.png")))

    assertions = {
        "模块核对清单": len(review) == 86,
        "截图合格": sum(row.screenshot_approved for row in review.values()) == 86,
        "红框合格": sum(row.annotation_approved for row in review.values()) == 86,
        "文字合格": sum(row.text_approved for row in review.values()) == 86,
        "标注配置": len(annotations) == 86,
        "标注图片": len(annotated_images) == 86,
        "流程节点": int(canvas["node_count"]) >= 4,
        "流程连接": int(canvas["connection_count"]) >= 3,
        "Word图注": captions == 86,
        "Word编号说明": legends == 86,
        "Word图片": len(document.inline_shapes) == 87,
        "PDF页数": pdf_pages == 142,
        "渲染页数": rendered_pages == pdf_pages,
    }
    failed = [name for name, passed in assertions.items() if not passed]
    print(
        f"MODULES={len(review)} SCREENSHOT_OK={sum(row.screenshot_approved for row in review.values())} "
        f"ANNOTATION_OK={sum(row.annotation_approved for row in review.values())} "
        f"TEXT_OK={sum(row.text_approved for row in review.values())} "
        f"NODES={canvas['node_count']} CONNECTIONS={canvas['connection_count']} "
        f"ANNOTATED_IMAGES={len(annotated_images)} CAPTIONS={captions} LEGENDS={legends} "
        f"INLINE_IMAGES={len(document.inline_shapes)} PDF_PAGES={pdf_pages} "
        f"RENDERED_PAGES={rendered_pages} ERRORS={len(failed)}"
    )
    for name in failed:
        print("ERROR=" + name)
    raise SystemExit(1 if failed else 0)


if __name__ == "__main__":
    main()
